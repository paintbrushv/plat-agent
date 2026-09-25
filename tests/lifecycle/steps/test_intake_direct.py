from __future__ import annotations

import json
import os
import subprocess
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from plat_agent.contracts.envelope import BridgeRequestV1
from plat_agent.lifecycle.state import LifecycleState
from plat_agent.lifecycle.steps.intake import (
    IntakeStep,
    _harvest_property_tax_evidence,
    _build_house_box_score_from_standardized,
    _classify_document_type,
    _enrich_canonical_supporting_docs,
    _normalize_rent_roll_with_market_study,
    _parse_box_score,
    _rebase_canonical_to_house_box_score,
    _select_document_by_type,
)


def test_intake_direct_subprocess_writes_outputs_without_llm_dispatch(tmp_path: Path) -> None:
    state = LifecycleState(deal_slug="willow_court", run_id="run_001")
    deal_root = tmp_path / "deals" / "willow_court"
    run_dir = deal_root / "outputs" / "run_001"
    raw_inputs = run_dir / "raw_inputs"
    raw_inputs.mkdir(parents=True)
    (raw_inputs / "Willow Court OM.pdf").write_bytes(b"%PDF")
    (raw_inputs / "Willow Court RR.xlsx").write_bytes(b"PK")
    (raw_inputs / "Willow Court T12.xlsx").write_bytes(b"PK")

    repo_root = tmp_path / "mfu"
    (repo_root / "runs").mkdir(parents=True)
    (repo_root / "runs" / "ingest_deal.py").write_text("# stub\n")

    captured = {}

    def fake_run(cmd, cwd=None, capture_output=False, text=False):
        cmd_str = " ".join(str(part) for part in cmd)
        if "pdftotext" in cmd_str:
            return subprocess.CompletedProcess(cmd, 0, stdout="", stderr="")
        if "standardize_rent_roll.py" in cmd_str:
            out_dir = Path(cmd[cmd.index("--output-dir") + 1])
            out_dir.mkdir(parents=True, exist_ok=True)
            (out_dir / "rent_roll_standardized.csv").write_text(
                "unit_id,floorplan_code,bed_type,bath_count,sqft,market_rent,lease_rent,status\n"
                "101,A1,1BR,1.0,700,1200,1100,Occupied\n"
            )
            (out_dir / "floorplan_summary.csv").write_text(
                "PlanCode,BedType,Units,SqFt,AvgMarketRent\n"
                "A1,1BR,1,700,1200\n"
            )
            return subprocess.CompletedProcess(cmd, 0, stdout="ok", stderr="")
        captured["cmd"] = list(cmd)
        captured["cwd"] = cwd
        output_path = Path(cmd[cmd.index("--output") + 1])
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(json.dumps({
            "metadata": {"analyst_review_required": True},
            "unit_cohorts": [{"cohort_id": "A1"}],
        }))
        return subprocess.CompletedProcess(cmd, 0, stdout="ok", stderr="")

    step = IntakeStep()
    request = step._build_request(state, run_dir=run_dir, deal_root=deal_root)

    with patch("plat_agent.lifecycle.steps.intake.subprocess.run", side_effect=fake_run):
        response = __import__("plat_agent.lifecycle.steps.intake", fromlist=["dispatch_sibling_agent"]).dispatch_sibling_agent(
            repo=type("Repo", (), {"path": repo_root, "name": "mfu"})(),
            request=request,
            payload_model=None,
        )

    assert response.status == "needs_analyst_input"
    assert "--rent-roll" in captured["cmd"]
    assert "--t12" in captured["cmd"]
    assert "--om" in captured["cmd"]
    assert (run_dir / "intake" / "canonical_deal.json").exists()
    assert (run_dir / "intake" / "manifest.md").exists()
    assert (run_dir / "intake" / "intake_punchlist.md").exists()


def test_intake_uses_standardized_fallback_when_rent_roll_source_is_pdf(tmp_path: Path) -> None:
    state = LifecycleState(deal_slug="huntington_hills", run_id="run_001")
    deal_root = tmp_path / "deals" / "huntington_hills"
    run_dir = deal_root / "outputs" / "run_001"
    raw_inputs = run_dir / "raw_inputs"
    raw_inputs.mkdir(parents=True)
    (raw_inputs / "Huntington Hills Rent Roll 4-30-26.pdf").write_bytes(b"%PDF")
    (raw_inputs / "Huntington Hills - Income Statment - April 2026.xlsx").write_bytes(b"PK")
    fallback = deal_root / "standardized" / "om_unit_mix_fallback_rent_roll.csv"
    fallback.parent.mkdir(parents=True)
    fallback.write_text("Unit,Type,Sqft,Beds,Baths,Status,Monthly Rent,Market Rent\n")

    repo_root = tmp_path / "mfu"
    (repo_root / "runs").mkdir(parents=True)
    (repo_root / "runs" / "ingest_deal.py").write_text("# stub\n")

    captured = {}

    def fake_run(cmd, cwd=None, capture_output=False, text=False):
        if "standardize_rent_roll.py" in " ".join(str(part) for part in cmd):
            captured["standardize_cmd"] = list(cmd)
            out_dir = Path(cmd[cmd.index("--output-dir") + 1])
            out_dir.mkdir(parents=True, exist_ok=True)
            (out_dir / "rent_roll_standardized.csv").write_text(
                "unit_id,floorplan_code,bed_type,bath_count,sqft,market_rent,lease_rent,status\n"
            )
            (out_dir / "floorplan_summary.csv").write_text("PlanCode,BedType,Units,SqFt,AvgMarketRent\n")
            return subprocess.CompletedProcess(cmd, 0, stdout="ok", stderr="")
        captured["cmd"] = list(cmd)
        output_path = Path(cmd[cmd.index("--output") + 1])
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(json.dumps({
            "metadata": {"analyst_review_required": True},
            "unit_cohorts": [{"cohort_id": "a1"}],
        }))
        return subprocess.CompletedProcess(cmd, 0, stdout="ok", stderr="")

    step = IntakeStep()
    request = step._build_request(state, run_dir=run_dir, deal_root=deal_root)

    with patch("plat_agent.lifecycle.steps.intake.subprocess.run", side_effect=fake_run):
        response = __import__("plat_agent.lifecycle.steps.intake", fromlist=["dispatch_sibling_agent"]).dispatch_sibling_agent(
            repo=type("Repo", (), {"path": repo_root, "name": "mfu"})(),
            request=request,
            payload_model=None,
        )

    assert response.status == "needs_analyst_input"
    assert captured["cmd"][captured["cmd"].index("--rent-roll") + 1] == str(fallback)
    assert str(fallback) in captured["standardize_cmd"]


def test_document_classifier_does_not_treat_commercial_as_om() -> None:
    assert _classify_document_type(Path("The Frank_03.23.2026 Commercial Rent Roll.pdf")) == "commercial_rent_roll"
    assert _classify_document_type(Path("WDIS_The Frank OM FINAL.pdf")) == "om"


def test_document_classifier_recognizes_debt_matrix_and_insurance_indication() -> None:
    assert (
        _classify_document_type(Path("IPADebtMatrix_Laurel Heights at Cityview_04.20.26.pdf"))
        == "debt_guidance"
    )
    assert (
        _classify_document_type(Path("Casa Nube Loan Information.pdf"))
        == "debt_guidance"
    )
    assert (
        _classify_document_type(Path("RPMF Indication - Laurel Heights at Cityview.pdf"))
        == "insurance_quote"
    )


def test_document_classifier_recognizes_profit_loss_recap_as_t12() -> None:
    assert (
        _classify_document_type(Path("Profit & Loss 12 Month Recap_20260416085311931.xlsx"))
        == "t12"
    )
    assert _classify_document_type(Path("Financials-Dylan-T12-4.2026.pdf")) == "t12"
    assert _classify_document_type(Path("Box-Score-Dylan-3.2026.pdf")) == "box_score"


def test_document_classifier_recognizes_broker_typos_and_pdf_rent_rolls() -> None:
    assert (
        _classify_document_type(Path("Huntington Hills - Income Statment - April 2026.xlsx"))
        == "t12"
    )
    assert (
        _classify_document_type(Path("Huntington Hills Rent Roll 4-30-26.pdf"))
        == "rent_roll"
    )


def test_document_classifier_recognizes_huntington_hills_optional_support_docs() -> None:
    assert _classify_document_type(Path("H Hills Improvements - 4.8.26.docx")) == "capex_quote"
    assert _classify_document_type(Path("Wythe - CapEx.xlsx")) == "capex_quote"
    assert _classify_document_type(Path("Capital expenses.docx")) == "capex_quote"
    assert _classify_document_type(Path("HH vacancy % - April 2026.xlsx")) == "occupancy_support"
    assert _classify_document_type(Path("JLL - Huntington Hills Flyer.pdf")) == "marketing_flyer"


def test_document_classifier_recognizes_supporting_diligence_docs() -> None:
    assert _classify_document_type(Path("Commitment - 9001312600146 .pdf")) == "title_doc"
    assert _classify_document_type(Path("Manitoba_Cantamar_Survey_20190708.PDF")) == "survey"
    assert _classify_document_type(Path("Valorem Report Villas at Cantamar 4.21.26.pdf")) == "tax_doc"
    assert _classify_document_type(Path("costar_norman_submarket.v1.pdf")) == "market_data"
    assert _classify_document_type(Path("Norman 7-Pack.pdf")) == "om"
    assert _classify_document_type(Path("source_onedrive_inventory.md")) == "source_inventory"
    assert _classify_document_type(Path("Villas at Cantamar - MSI_LossRun 25-26.pdf")) == "insurance_quote"
    assert _classify_document_type(Path("Villas at Cantamar CM Quote 4-17-26.pdf")) == "debt_guidance"
    assert _classify_document_type(Path("Loan Details - Villas At Cantamar Apartments.xlsx")) == "existing_debt"
    assert _classify_document_type(Path("Statement_19-43423_2026-04-01.PDF")) == "existing_debt"
    assert _classify_document_type(Path("Villas at Cantamar SMBC SOFR Replacement Rate Cap Agreement (2024).pdf")) == "existing_debt"
    assert _classify_document_type(Path("PB HOLDINGS IV LLC 44MO CONTRACT Electricity.pdf")) == "service_contract"
    assert _classify_document_type(Path("Republic Services_TrashPickupServices.pdf")) == "service_contract"
    assert _classify_document_type(Path("Grant at Valley Ranch - AR - March 2026.xlsx")) == "delinquency_report"
    assert _classify_document_type(Path("VaC spectrum contract.pdf")) == "service_contract"
    assert (
        _classify_document_type(Path("Villas of Cantamar Chiller Maintenance Agreement Aug - Dec 2023_signed.pdf"))
        == "service_contract"
    )
    assert _classify_document_type(Path("Villas at Cantamar Deleinquency 3-31-26.pdf")) == "delinquency_report"


def test_parse_resman_pdf_box_score(monkeypatch, tmp_path: Path) -> None:
    path = tmp_path / "Box-Score-Dylan-3.2026.pdf"
    path.write_bytes(b"%PDF")
    monkeypatch.setattr(
        "plat_agent.lifecycle.steps.intake._load_pdf_layout_text",
        lambda pdf: """
Box Score
Occupancy
Unit     Total    Excl.                        Net                                           Pre- Vacant Not Vacant
Type     Units    Units Total M/I Total M/O Change             Occ     % Occ Vacant        Leased    Leased Ready
A1-1X1     16         0        0          0      0              16    100.0%      0             0          0      0
B1-2X1     57         0        3          1      2              53     93.0%      4             1          3      2
B1U -2x1    1         0        0          0      0               1    100.0%      0             0          0      0
B2-2X2     16         0        0          0      0              16    100.0%      0             0          0      0
C1-3X2     32         0        2          1      1              32    100.0%      0             0          0      0
Total     122         0        5          2      3             118     96.7%      4             1          3      2
""",
    )

    parsed = _parse_box_score(path)

    assert parsed is not None
    assert parsed["total_units"] == 122
    by_code = {row["code"]: row for row in parsed["floorplans"]}
    assert by_code["B1-2X1"]["units"] == 57
    assert by_code["B1-2X1"]["available_units"] == 4


def test_enrichment_rebases_empty_local_cohorts_from_standardized_roll(monkeypatch, tmp_path: Path) -> None:
    rent_roll = tmp_path / "Rent-Roll-Dylan.xlsx"
    rent_roll.write_bytes(b"PK")
    monkeypatch.setattr(
        "plat_agent.lifecycle.steps.intake._normalize_rent_roll_with_market_study",
        lambda path, output_dir: {
            "status": "ok",
            "parser": "parse_vertical_charges",
            "standardized_rows": [
                {"floorplan_code": "A1-1X1", "status": "Occupied", "lease_rent": "1200", "market_rent": "1200"},
            ],
            "summary_rows": [
                {"PlanCode": "A1-1X1", "BedType": "1BR", "Units": "16", "SqFt": "554", "AvgMarketRent": "1200"},
            ],
        },
    )
    canonical = {"metadata": {}, "unit_cohorts": []}

    enriched = _enrich_canonical_supporting_docs(
        canonical,
        classified_single={"rent_roll": rent_roll},
        start_period="2026-06",
        end_period="2031-05",
        intake_dir=tmp_path,
    )

    assert enriched["unit_cohorts"][0]["unit_count"] == 16
    assert enriched["market_rent_curve"][0]["start_period"] == "2026-06"


def test_enrichment_uses_box_score_counts_with_standardized_rents(monkeypatch, tmp_path: Path) -> None:
    rent_roll = tmp_path / "Rent-Roll-Dylan.xlsx"
    rent_roll.write_bytes(b"PK")
    box_score = tmp_path / "Box-Score-Dylan.pdf"
    box_score.write_bytes(b"%PDF")
    monkeypatch.setattr(
        "plat_agent.lifecycle.steps.intake._normalize_rent_roll_with_market_study",
        lambda path, output_dir: {
            "status": "ok",
            "parser": "parse_vertical_charges",
            "standardized_rows": [
                {"floorplan_code": "B1-2X1", "status": "Occupied", "lease_rent": "1300", "market_rent": "1300"},
            ],
            "summary_rows": [
                {"PlanCode": "B1-2X1", "BedType": "2BR", "Units": "51", "SqFt": "781", "AvgMarketRent": "1300"},
            ],
        },
    )
    monkeypatch.setattr(
        "plat_agent.lifecycle.steps.intake._parse_box_score",
        lambda path: {
            "source_file": "Box-Score-Dylan.pdf",
            "floorplans": [{"code": "B1-2X1", "units": 57, "available_units": 4}],
            "total_units": 57,
        },
    )
    canonical = {"metadata": {}, "unit_cohorts": []}

    enriched = _enrich_canonical_supporting_docs(
        canonical,
        classified_single={"rent_roll": rent_roll, "box_score": box_score},
        start_period="2026-06",
        end_period="2031-05",
        intake_dir=tmp_path,
    )

    assert enriched["metadata"]["property_summary"]["house_box_score"]["total_units"] == 57
    assert enriched["unit_cohorts"][0]["unit_count"] == 57
    assert enriched["unit_cohorts"][0]["initial_inplace_rent"] == 1300


def test_build_house_box_score_from_standardized_carries_bath_count() -> None:
    house_box = _build_house_box_score_from_standardized(
        summary_rows=[
            {"PlanCode": "B2", "BedType": "2BR", "Units": "3", "SqFt": "950", "AvgMarketRent": "1700"},
        ],
        standardized_rows=[
            {
                "floorplan_code": "B2",
                "status": "Occupied",
                "lease_rent": "1600",
                "market_rent": "1750",
                "bath_count": "1.5",
            },
            {
                "floorplan_code": "B2",
                "status": "Vacant",
                "lease_rent": "",
                "market_rent": "1850",
                "bath_count": "1.5",
            },
        ],
    )

    assert house_box is not None
    floorplan = house_box["floorplans"][0]
    assert floorplan["code"] == "B2"
    assert floorplan["units"] == 3
    assert floorplan["bedrooms"] == 2
    assert floorplan["bathrooms"] == 1.5
    assert floorplan["avg_in_place_rent"] == 1600
    assert floorplan["avg_market_rent"] == 1800


def test_rebase_canonical_replaces_local_rents_even_when_unit_count_matches() -> None:
    canonical = {
        "metadata": {"intake_sanity_flags": ["target_monthly_rent_missing_for_cohort_0br1ba"]},
        "unit_cohorts": [
            {
                "cohort_id": "a1",
                "unit_type": "A1",
                "unit_count": 10,
                "initial_inplace_rent": 999,
                "target_monthly_rent": 999,
                "target_monthly_rent_source": "local_parser",
            },
        ],
        "market_rent_curve": [
            {
                "cohort_id": "a1",
                "start_period": "2026-01",
                "end_period": "2030-12",
                "market_rent": 999,
            },
        ],
    }
    house_box = {
        "floorplans": [
            {
                "code": "A1",
                "units": 10,
                "avg_in_place_rent": 1210,
                "avg_market_rent": 1325,
                "avg_sqft": 650,
                "bedrooms": 1,
                "bathrooms": 1.0,
            },
        ],
        "total_units": 10,
        "status_counts": {"Occupied": 9, "Vacant": 1},
    }

    _rebase_canonical_to_house_box_score(
        canonical,
        house_box,
        start_period="2026-06",
        end_period="2031-05",
    )

    assert canonical["unit_cohorts"] == [
        {
            "cohort_id": "a1",
            "unit_type": "A1",
            "unit_count": 10,
            "initial_inplace_rent": 1210.0,
            "target_monthly_rent": 1325.0,
            "target_monthly_rent_source": "rent_roll_avg",
            "sqft": 650.0,
            "bedrooms": 1,
            "bathrooms": 1.0,
        },
    ]
    assert canonical["market_rent_curve"] == [
        {
            "cohort_id": "a1",
            "start_period": "2026-06",
            "end_period": "2031-05",
            "market_rent": 1325.0,
        },
    ]
    assert canonical["physical_vacancy_curve"][0]["vacancy_rate"] == 0.1
    assert canonical["metadata"]["intake_sanity_flags"] == ["unit_cohorts_rebased_to_standardized_rent_roll"]


def test_rebase_canonical_counts_lowercase_vacant_and_non_revenue_statuses() -> None:
    canonical = {
        "metadata": {},
        "unit_cohorts": [
            {
                "cohort_id": "a1",
                "unit_type": "A1",
            },
        ],
    }
    house_box = {
        "floorplans": [
            {
                "code": "A1",
                "units": 10,
                "avg_in_place_rent": 1210,
                "avg_market_rent": 1325,
            },
        ],
        "total_units": 10,
        "status_counts": {"Occupied": 7, "vacant": 1, "non revenue": 2},
    }

    _rebase_canonical_to_house_box_score(
        canonical,
        house_box,
        start_period="2026-06",
        end_period="2031-05",
    )

    assert canonical["physical_vacancy_curve"] == [
        {
            "cohort_id": "a1",
            "start_period": "2026-06",
            "end_period": "2031-05",
            "vacancy_rate": 0.3,
        },
    ]

def test_select_document_by_type_prefers_latest_dated_rent_roll() -> None:
    selected = _select_document_by_type({
        "rent_roll": [
            Path("Casa Nube Rent Roll - 4.13.26.xls"),
            Path("Casa Nube Rent Roll - 5.7.26.xls"),
        ],
        "t12": [Path("Casa Nube PL T12 March 30 2026.xlsx")],
    })

    assert selected["rent_roll"].name == "Casa Nube Rent Roll - 5.7.26.xls"


def test_select_document_by_type_prefers_cleaned_copy_for_same_date(
    tmp_path: Path,
) -> None:
    cleaned = tmp_path / "GVR RR 7.8.26 - SH CLEANED.xls"
    raw = tmp_path / "GVR RR 7.8.26.xls"
    cleaned.write_bytes(b"cleaned")
    raw.write_bytes(b"raw")
    cleaned.touch()
    raw.touch()

    selected = _select_document_by_type({"rent_roll": [cleaned, raw]})

    assert selected["rent_roll"] == cleaned


def test_select_document_by_type_prefers_newer_undated_copy_by_mtime(
    tmp_path: Path,
) -> None:
    cleaned = tmp_path / "Rent Roll CLEANED.xlsx"
    raw = tmp_path / "Rent Roll.xlsx"
    cleaned.write_bytes(b"cleaned")
    raw.write_bytes(b"raw")
    os.utime(cleaned, (100, 100))
    os.utime(raw, (200, 200))

    selected = _select_document_by_type({"rent_roll": [cleaned, raw]})

    assert selected["rent_roll"] == raw


def test_market_study_standardizer_uses_deal_specific_mapping(
    monkeypatch,
    tmp_path: Path,
) -> None:
    repo = tmp_path / "market-study-agent"
    (repo / "etl").mkdir(parents=True)
    (repo / "etl" / "standardize_rent_roll.py").write_text("# stub\n")
    mappings = repo / "configs" / "rent_roll_mappings"
    mappings.mkdir(parents=True)
    (mappings / "gardens_of_valley_ranch.yaml").write_text(
        "property_id: the_gardens_of_valley_ranch\n"
    )
    rent_roll = (
        tmp_path
        / "runs"
        / "deals"
        / "the_gardens_of_valley_ranch"
        / "outputs"
        / "run_002"
        / "raw_inputs"
        / "GVR RR 7.8.26 - SH CLEANED.xls"
    )
    rent_roll.parent.mkdir(parents=True)
    rent_roll.write_bytes(b"source")
    commands: list[list[str]] = []

    monkeypatch.setattr(
        "plat_agent.lifecycle.steps.intake.SiblingRepo.from_env_or_default",
        lambda *args, **kwargs: type("Repo", (), {"path": repo})(),
    )

    def fail_after_capture(command, **kwargs):
        commands.append(command)
        return subprocess.CompletedProcess(command, 1, "", "stop")

    monkeypatch.setattr(
        "plat_agent.lifecycle.steps.intake.subprocess.run",
        fail_after_capture,
    )

    _normalize_rent_roll_with_market_study(
        rent_roll,
        output_dir=tmp_path / "clean",
    )

    assert "--config" in commands[0]
    assert commands[0][commands[0].index("--config") + 1] == "gardens_of_valley_ranch"
    assert "--auto-detect" not in commands[0]


def test_intake_direct_classifies_and_enriches_box_score_and_debt_guidance(
    tmp_path: Path,
) -> None:
    state = LifecycleState(deal_slug="maple_grove", run_id="run_001")
    deal_root = tmp_path / "deals" / "maple_grove"
    run_dir = deal_root / "outputs" / "run_001"
    raw_inputs = run_dir / "raw_inputs"
    raw_inputs.mkdir(parents=True)
    (raw_inputs / "The Demo at Maple Grove OM.pdf").write_bytes(b"%PDF")
    (raw_inputs / "Maple Grove Rent Roll 4.14.26.xlsx").write_bytes(b"PK")
    (raw_inputs / "Maple Grove T12 March.xlsx").write_bytes(b"PK")
    (raw_inputs / "WD terms.pdf").write_bytes(b"%PDF")

    box_score = raw_inputs / "ops_snapshot.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "BoxScore Summary"
    ws["A2"] = "Maple Grove Apartments (briarc)"
    ws["A3"] = "Date = 06/30/2022-04/14/2026"
    ws.append(["Availability"])
    ws.append(["Code", "Name", "Avg. Sq Ft.", "Avg. Rent", "Units", "Occupied No Notice", "Vacant Rented", "Vacant Unrented", "Notice Rented", "Notice Unrented", "Avail", "Model"])
    ws.append(["bc_B1", "2x2 925 sqft", "925", "1,847", "100", "86", "2", "7", "0", "4", "11", "1"])
    ws.append(["Total", "960", "1,867", "238", "206", "3", "16", "1", "11", "27", "1", "0"])
    wb.save(box_score)

    repo_root = tmp_path / "mfu"
    (repo_root / "runs").mkdir(parents=True)
    (repo_root / "runs" / "ingest_deal.py").write_text("# stub\n")

    debt_guidance_text = """
    PROPOSED FINANCING TERMS
    Loan Option
    Agency 5 Year Fixed
    Fixed
    $24,806,000
    Max LTV
    70.0%
    Min DSCR
    1.25x
    U W NOI
    $1,976,327
    Benchmark
    5 Year Treasury
    Benchmark Rate
    3.91%
    Spread
    1.50% - 1.65%
    All-in-Rate
    5.41% - 5.56%
    """
    om_text = """
    PROPERTY SUMMARY
    A DDRES S
    742 Sample Ave, Demo City, ST 00000
    FOO
    Y E A R BUILT
    1984
    TOTAL UNITS
    238

    INCOME:
    Gross Scheduled Market Rent
    Plus: Renovated Unit Premiums
    Gain / Loss to Lease
    Gross Potential Rent

    AMOUNT
    Amount
    $5,355,600
    $0
    ($891,636)
    $4,463,964

    T3 REVENUE/T12 EXPENSES
    % OF GPR
    GPR
    120%
    0%
    -20.0%
    100.0%

    PER UNIT
    Unit
    $22,503
    $0
    ($3,746)
    $18,756

    AMOUNT
    Amount
    $4,354,709
    $87,267
    $15,697
    $4,457,673

    PROFORMA
    % OF GPR
    GPR
    98%
    2%
    0.4%
    100.0%

    PER UNIT
    Unit
    $18,297
    $367
    $66
    $18,730

    Less: Vacancy
    Less: Concessions
    Less: Non-Revenue Units
    Less: Bad Debt
    Net Rental Income

    ($392,067)
    ($38,953)
    ($30,918)
    ($30,382)
    $3,971,643

    -8.8%
    -0.9%
    -0.7%
    -0.7%
    89.0%

    ($1,647)
    ($164)
    ($130)
    ($128)
    $16,688

    ($222,884)
    ($44,577)
    ($28,095)
    ($22,288)
    $4,139,830

    -5.0%
    -1.0%
    -0.6%
    -0.5%
    92.9%

    ($936)
    ($187)
    ($118)
    ($94)
    $17,394

    Plus: Utility Reimbursements
    Plus: Parking/Storage Income
    Plus: Internet Package
    Plus: Other Income
    Total Income
    Monthly Collections

    $305,776
    $8,703
    $0
    $143,367
    $4,429,489
    $369,124

    6.8%
    0.2%
    0.0%
    3.2%
    99.2%

    $1,285
    $37
    $0
    $602
    $18,611

    $319,634
    $10,080
    $78,540
    $192,098
    $4,740,182
    $395,015

    7.2%
    0.2%
    1.8%
    4.3%
    106.3%

    $1,343
    $42
    $330
    $807
    $19,917

    EXPENSES:
    Utilities
    Repairs and Maintenance
    Apartment Make Ready
    Contract Services
    Marketing
    Payroll Expenses
    General and Administrative
    Property Taxes
    Franchise Tax
    Insurance
    Management Fees
    Miscellaneous
    Total Expenses

    $299,219
    $160,936
    $98,493
    $176,012
    $49,122
    $498,451
    $105,791
    $752,437
    $14,951
    $298,787
    $190,125
    $0
    $2,644,324

    6.7%
    3.6%
    2.2%
    3.9%
    1.1%
    11.2%
    2.4%
    16.9%
    0.3%
    6.7%
    4.3%
    0.0%
    59.2%

    $1,257
    $676
    $414
    $740
    $206
    $2,094
    $444
    $3,162
    $63
    $1,255
    $799
    $0
    $11,111

    $299,219
    $119,000
    $83,300
    $176,012
    $47,600
    $403,000
    $105,791
    $752,437
    $15,690
    $166,600
    $142,205
    $0
    $2,310,854

    6.7%
    2.7%
    1.9%
    3.9%
    1.1%
    9.0%
    2.4%
    16.9%
    0.4%
    3.7%
    3.2%
    0.0%
    51.8%

    $1,257
    $500
    $350
    $740
    $200
    $1,693
    $444
    $3,162
    $66
    $700
    $598
    $0
    $9,709

    Net Operating Income
    Replacement Reserves
    Net Cash Flow

    $1,785,165
    $0
    $1,785,165

    40.0%
    0.0%
    40.0%

    $7,501
    $0
    $7,501

    $2,429,328
    ($59,500)
    $2,369,828

    54.5%
    -1.3%
    53.2%

    $10,207
    ($250)
    $9,957

    PROPERTY TAXES
    TAXING JURISDICTION
    TAX RATE PER $100
    Sample City
    0.5375
    Sample ISD
    0.9481
    Sample County
    0.2155
    Sample College
    0.106575
    Sample Hospital
    0.212
    *Per SampleCAD
    0.02019675
    **Property Account #: 14006580000000000
    2026 Tax Rate
    The WDIS Proforma assumes $1,693/unit for Payroll.
    The WDIS Proforma assumes Replacement Reserves of $250/unit.
    """

    def fake_run(cmd, cwd=None, capture_output=True, text=True):
        if cmd and cmd[0] == "pdftotext":
            source = str(cmd[1])
            stdout = debt_guidance_text if "WD terms" in source else om_text
            return subprocess.CompletedProcess(cmd, 0, stdout=stdout, stderr="")
        cmd_str = " ".join(str(part) for part in cmd)
        if "standardize_rent_roll.py" in cmd_str:
            out_dir = Path(cmd[cmd.index("--output-dir") + 1])
            out_dir.mkdir(parents=True, exist_ok=True)
            (out_dir / "rent_roll_standardized.csv").write_text(
                "unit_id,floorplan_code,bed_type,bath_count,sqft,market_rent,lease_rent,status\n"
                "0101,bc_B1,2BR,2.0,925,1847,1547,Occupied\n"
                "0102,bc_B1,2BR,2.0,925,1847,,Vacant\n"
                "0103,bc_B1,2BR,2.0,925,1847,1547,Occupied\n"
            )
            (out_dir / "floorplan_summary.csv").write_text(
                "PlanCode,BedType,Units,SqFt,AvgMarketRent\n"
                "bc_B1,2BR,238,925,1847\n"
            )
            return subprocess.CompletedProcess(cmd, 0, stdout="ok", stderr="")
        output_path = Path(cmd[cmd.index("--output") + 1])
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(json.dumps({
            "metadata": {
                "address": "742 Sample Ave, Demo City, ST 00000",
                "intake_sanity_flags": [
                    "om_extraction_unavailable",
                    "standardized_rent_roll_unit_count_mismatch",
                    "box_score_unit_count_mismatch",
                ],
            },
            "time_grid": {"analysis_start_date": "2026-06-01", "analysis_end_date": "2031-05-01"},
            "unit_cohorts": [{"cohort_id": "bc_b1", "unit_count": 242}],
        }))
        return subprocess.CompletedProcess(cmd, 0, stdout="ok", stderr="")

    step = IntakeStep()
    request = step._build_request(state, run_dir=run_dir, deal_root=deal_root)

    with patch("plat_agent.lifecycle.steps.intake.subprocess.run", side_effect=fake_run):
        response = __import__("plat_agent.lifecycle.steps.intake", fromlist=["dispatch_sibling_agent"]).dispatch_sibling_agent(
            repo=type("Repo", (), {"path": repo_root, "name": "mfu"})(),
            request=request,
            payload_model=None,
        )

    assert response.status == "ok"
    payload = response.payload
    assert payload["documents_classified"]["ops_snapshot.xlsx"] == "box_score"
    assert payload["documents_classified"]["WD terms.pdf"] == "debt_guidance"

    canonical = json.loads((run_dir / "intake" / "canonical_deal.json").read_text())
    property_summary = canonical["metadata"]["property_summary"]
    assert property_summary["house_box_score"]["total_units"] == 238
    assert property_summary["house_box_score"]["status_counts"]["Occupied"] == 2
    assert property_summary["house_box_score"]["floorplans"][0]["avg_in_place_rent"] == 1547.0
    assert property_summary["house_box_score"]["floorplans"][0]["avg_market_rent"] == 1847.0
    assert canonical["unit_cohorts"][0]["target_monthly_rent_source"] == "rent_roll_avg"
    assert property_summary["box_score"]["total_units"] == 238
    assert property_summary["debt_guidance"]["uw_noi"] == 1_976_327.0
    assert property_summary["debt_guidance"]["hold_matched_recommendation"]["benchmark_rate"] == 0.0391
    assert canonical["metadata"]["year_built"] == 1984
    broker_snapshot = property_summary["broker_underwriting_snapshot"]
    assert canonical["metadata"]["address"] == "742 Sample Ave, Demo City, ST 00000"
    assert broker_snapshot["unit_count"] == 238
    assert broker_snapshot["address"] == "742 Sample Ave, Demo City, ST 00000"
    assert broker_snapshot["noi"] == 2_429_328.0
    assert broker_snapshot["expense_assumptions"]["insurance"] == 166_600.0
    assert broker_snapshot["property_tax_context"]["local_tax_rate"] == 0.02019675
    assert property_summary["property_tax_evidence_candidates"] == [
        {
            "millage_rate_mills": 20.19675,
            "source": "offering_memorandum",
            "source_locator": (
                "raw_inputs/The Demo at Maple Grove OM.pdf, page 1, "
                "line 277, table PROPERTY TAXES, label TAX RATE PER $100"
            ),
        }
    ]
    assert property_summary["property_tax_evidence_inspected_paths"] == [
        "raw_inputs/The Demo at Maple Grove OM.pdf",
        "raw_inputs/WD terms.pdf",
    ]
    assert broker_snapshot["parser_metadata"]["parser_family"] == "generic_table_om_v1"
    assert "property_tax_section" in broker_snapshot["parser_metadata"]["matched_sections"]
    assert "box_score_unit_count_mismatch" in canonical["metadata"]["intake_sanity_flags"]
    assert "standardized_rent_roll_unit_count_mismatch" in canonical["metadata"]["intake_sanity_flags"]
    punchlist = (run_dir / "intake" / "intake_punchlist.md").read_text()
    assert "No blocking intake issues detected." in punchlist
    assert "resolved by consensus" in punchlist
    assert "deterministic OM harvesting captured core broker underwriting facts" in punchlist


def test_harvest_property_tax_evidence_inspects_all_supported_documents(
    tmp_path: Path,
) -> None:
    raw = tmp_path / "raw_inputs"
    raw.mkdir()
    om = raw / "OM.pdf"
    notice = raw / "notice.pdf"
    ambiguous = raw / "ambiguous-tax.pdf"
    debt = raw / "lender-terms.pdf"
    for path in (om, notice, ambiguous, debt):
        path.write_bytes(b"%PDF")
    texts = {
        om: "Combined property tax rate: 2.49%",
        notice: "Combined ad valorem tax rate: 2.531 per $100",
        ambiguous: "Combined property tax millage: see assessor schedule",
        debt: "Aggregate property tax rate: 0.02500 decimal rate",
    }
    canonical = {"metadata": {}}

    with patch(
        "plat_agent.lifecycle.steps.intake._shared_load_pdf_text",
        side_effect=lambda path: texts[path],
    ):
        enriched = _harvest_property_tax_evidence(
            canonical,
            classified={
                "om": [om],
                "tax_doc": [notice, ambiguous],
                "debt_guidance": [debt],
            },
            raw_inputs_dir=raw,
        )

    summary = enriched["metadata"]["property_summary"]
    assert summary["property_tax_evidence_candidates"] == [
        {
            "millage_rate_mills": 24.9,
            "source": "offering_memorandum",
            "source_locator": (
                "raw_inputs/OM.pdf, page 1, line 1, "
                "label Combined property tax rate"
            ),
        },
        {
            "millage_rate_mills": 25.31,
            "source": "county_tax_notice",
            "source_locator": (
                "raw_inputs/notice.pdf, page 1, line 1, "
                "label Combined ad valorem tax rate"
            ),
        },
        {
            "millage_rate_mills": 25.0,
            "source": "debt_guidance",
            "source_locator": (
                "raw_inputs/lender-terms.pdf, page 1, line 1, "
                "label Aggregate property tax rate"
            ),
        },
    ]
    assert summary["property_tax_evidence_inspected_paths"] == [
        "raw_inputs/OM.pdf",
        "raw_inputs/ambiguous-tax.pdf",
        "raw_inputs/lender-terms.pdf",
        "raw_inputs/notice.pdf",
    ]
    assert summary["property_tax_evidence_locations"] == [
        "raw_inputs/OM.pdf, page 1, line 1, label Combined property tax rate",
        "raw_inputs/ambiguous-tax.pdf, page 1, line 1, "
        "label Combined property tax millage",
        "raw_inputs/lender-terms.pdf, page 1, line 1, "
        "label Aggregate property tax rate",
        "raw_inputs/notice.pdf, page 1, line 1, "
        "label Combined ad valorem tax rate",
    ]
